from docx import Document
from docx.shared import Pt
from pathlib import Path

import sys

base = Path(__file__).parent
stem = sys.argv[1] if len(sys.argv) > 1 else "PRD-Platform-Requirements-and-Specification-v1"
md_path = base / f"{stem}.md"
out_path = base / f"{stem}.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

lines = md_path.read_text(encoding="utf-8").splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if line.startswith("# "):
        doc.add_heading(line[2:].strip(), level=0)
    elif line.startswith("## "):
        doc.add_heading(line[3:].strip(), level=1)
    elif line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=2)
    elif line.strip() == "---":
        pass
    elif line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            row = [c.strip() for c in lines[i].split("|")[1:-1]]
            if row and not all(set(c) <= set("- :") for c in row):
                rows.append(row)
            i += 1
        i -= 1
        if rows:
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    table.rows[ri].cells[ci].text = cell
    elif line.strip().startswith("- "):
        doc.add_paragraph(line.strip()[2:], style="List Bullet")
    elif line.strip():
        doc.add_paragraph(line.strip())
    i += 1

doc.save(out_path)
print(f"Wrote {out_path}")
