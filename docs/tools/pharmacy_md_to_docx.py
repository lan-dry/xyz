#!/usr/bin/env python3
"""Convert pharmacy project markdown specs to Word (.docx) for client signing."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt
def set_document_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(6)


def add_rich_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    text = text.strip()
    if not text:
        return
    p = doc.add_paragraph(style=style)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)


def parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_table_separator(line: str) -> bool:
    s = line.strip().replace("|", "").replace(":", "").replace("-", "").strip()
    return len(s) == 0 and "|" in line and "-" in line


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(col_count):
            cell_text = row[j] if j < len(row) else ""
            cell_text = re.sub(r"\*\*([^*]+)\*\*", r"\1", cell_text)
            table.rows[i].cells[j].text = cell_text
    doc.add_paragraph()


def convert_md_to_docx(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_document_defaults(doc)

    i = 0
    in_code = False
    table_buffer: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            add_table(doc, table_buffer)
            table_buffer = []

    while i < len(lines):
        line = lines[i]
        raw = line

        if line.strip().startswith("```"):
            in_code = not in_code
            i += 1
            continue

        if in_code:
            p = doc.add_paragraph(style="Intense Quote")
            p.add_run(line)
            i += 1
            continue

        if is_table_separator(line):
            i += 1
            continue

        if "|" in line and line.strip().startswith("|"):
            if is_table_separator(line):
                i += 1
                continue
            table_buffer.append(parse_table_row(line))
            i += 1
            continue

        flush_table()

        stripped = line.strip()

        if stripped in ("---", "***", "___"):
            doc.add_paragraph()
            i += 1
            continue

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
            i += 1
            continue
        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=4)
            i += 1
            continue

        if stripped.startswith("- [ ]") or stripped.startswith("- [x]"):
            add_rich_paragraph(doc, stripped[2:].strip(), style="List Bullet")
            i += 1
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            add_rich_paragraph(doc, stripped[2:].strip(), style="List Bullet")
            i += 1
            continue

        num_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if num_match:
            add_rich_paragraph(doc, num_match.group(2), style="List Number")
            i += 1
            continue

        if stripped.startswith(">"):
            add_rich_paragraph(doc, stripped.lstrip("> ").strip(), style="Intense Quote")
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        add_rich_paragraph(doc, stripped)
        i += 1

    flush_table()
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    out_dir = root / "pharmacy-platform" / "word"
    files = [
        ("PHARMACY-ONLINE-PLATFORM-SRS.md", "SAL-PHARM-SRS-001-v2.0-Online-Pharmacy-SRS.docx"),
        ("PHARMACY-ONLINE-PLATFORM-QUOTE.md", "SAL-PHARM-QT-001-Online-Pharmacy-Quotation.docx"),
        ("PHARMACY-ONLINE-PLATFORM-SOW.md", "SAL-PHARM-SOW-001-Online-Pharmacy-Agreement.docx"),
        ("PHARMACY-ONLINE-ORDER-PLATFORM-SPEC.md", "PHARMACY-ONLINE-ORDER-PLATFORM-SPEC-Superseded.docx"),
    ]
    for md_name, docx_name in files:
        md_path = root / md_name
        if not md_path.exists():
            print(f"Skip missing: {md_path}", file=sys.stderr)
            continue
        docx_path = out_dir / docx_name
        convert_md_to_docx(md_path, docx_path)
        print(f"Wrote {docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
