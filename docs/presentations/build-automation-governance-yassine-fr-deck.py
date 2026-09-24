"""Generate French Salanor Automation & Governance deck (.pptx).

Run: python docs/presentations/build-automation-governance-yassine-fr-deck.py

Output:
  docs/presentations/Salanor-Automation-Gouvernance.pptx
  e:/salanor/Presentation/Yassine/Salanor-Automation-Gouvernance.pptx
"""

from __future__ import annotations

import importlib.util
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("deck_base", ROOT / "build-automation-governance-deck.py")
d = importlib.util.module_from_spec(spec)
assert spec.loader
spec.loader.exec_module(d)

OUT_REPO = ROOT / "Salanor-Automation-Gouvernance.pptx"
OUT_YASSINE = Path("e:/salanor/Presentation/Yassine/Salanor-Automation-Gouvernance.pptx")


def slide_title(prs) -> None:
    slide = d.blank(prs)
    d.add_rect(slide, d.Inches(0), d.Inches(0), d.Inches(10), d.Inches(0.06), d.TEAL_DIM)
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(0.55),
        d.CONTENT_W,
        d.Inches(0.28),
        "SALANOR · CONFIDENTIEL · 2026",
        size=9,
        color=d.TEAL,
        bold=True,
    )
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(1.15),
        d.Inches(8.5),
        d.Inches(1.4),
        "Automatisation et gouvernance\npour la finance",
        size=36,
        bold=True,
    )
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(2.75),
        d.Inches(8.2),
        d.Inches(1.0),
        "Nous automatisons vos flux critiques et intégrons Aegis dès le premier jour : "
        "vitesse, contrôle et preuve pour paiements, sinistres et KYC.",
        size=15,
        color=d.MUTED,
    )
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(6.35),
        d.CONTENT_W,
        d.Inches(0.35),
        "Landry Bougang Fotso · Fondateur, Salanor Ltd · Kigali",
        size=11,
        color=d.DIM,
    )
    d.add_notes(
        slide,
        "Présentez-vous. Histoire en trois actes : pourquoi automatiser, le risque, Salanor + Aegis. "
        "Banque et assurance uniquement.",
    )


def slide_reality(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Acte 1 · L'opportunité", "Beaucoup d'institutions gèrent encore l'essentiel à la main")
    d.write_bullets(
        slide,
        d.MARGIN_L,
        d.Inches(2.35),
        d.CONTENT_W,
        d.Inches(3.5),
        [
            "Chaînes d'e-mails, tableurs, transmissions papier, saisies en double",
            "Délais longs : les clients attendent, les équipes relancent les approbations",
            "Décisions incohérentes selon qui est de service",
            "Erreurs qui pèsent davantage quand le volume augmente",
            "Le savoir-processus part avec les personnes qui quittent l'entreprise",
        ],
        size=13,
    )
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(5.95),
        d.CONTENT_W,
        d.Inches(0.55),
        "Les attentes numériques montent partout. Les processus manuels deviennent le goulot d'étranglement.",
        size=12,
        color=d.MUTED,
    )
    d.slide_footer(slide, "Automatisation et gouvernance · Finance et assurance")
    d.add_notes(
        slide,
        "Pas réservé à l'Afrique : beaucoup d'institutions régulées restent manuelles. "
        "Transition : un seul flux bien fait suffit pour commencer.",
    )


def slide_benefits(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(
        slide,
        "Ce que l'automatisation apporte",
        "Automatisation structurée de flux, pas « remplacer tout le monde par l'IA »",
    )
    cards = [
        ("Vitesse", "Les étapes routinières en minutes, pas en heures ou jours"),
        ("Cohérence", "Même règle, même chemin, à chaque fois"),
        ("Échelle", "Plus de volume sans croissance linéaire des effectifs"),
        ("Focus", "Les équipes sur le jugement, pas sur le copier-coller"),
    ]
    w = d.Inches(2.05)
    gap = d.Inches(0.18)
    y = d.Inches(2.45)
    h = d.Inches(2.2)
    for i, (head, body) in enumerate(cards):
        x = d.MARGIN_L + i * (w + gap)
        d.add_round_card(slide, x, y, w, h)
        d.write_text(slide, x + d.Inches(0.15), y + d.Inches(0.18), w - d.Inches(0.25), d.Inches(0.35), head, size=13, bold=True, color=d.TEAL)
        d.write_text(slide, x + d.Inches(0.15), y + d.Inches(0.58), w - d.Inches(0.25), d.Inches(1.4), body, size=11, color=d.MUTED)
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(5.0),
        d.CONTENT_W,
        d.Inches(0.5),
        "L'automatisation retire la friction autour des décisions sensibles. Elle ne retire pas les humains de la décision.",
        size=12,
        color=d.MUTED,
    )
    d.add_notes(slide, "Définir l'automatisation comme règles + routage + notifications. Éviter le hype IA.")


def slide_banking(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Banque", "Points de départ à forte valeur. Choisissez un flux.")
    d.write_bullets(
        slide,
        d.MARGIN_L,
        d.Inches(2.35),
        d.CONTENT_W,
        d.Inches(3.2),
        [
            "Virements sortants au-dessus d'un seuil : approbation avant exécution",
            "Contrôles KYC / LBC-FT : router les exceptions, signaler les dossiers incomplets",
            "Changements de limite ou de statut de compte : double contrôle avant mise à jour du core",
            "Alertes de rapprochement : remonter les écarts avec contexte à la bonne équipe",
        ],
        size=13,
    )
    d.add_notes(slide, "Choisissez le flux qui empêche ops ou conformité de dormir la nuit.")


def slide_insurance(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Assurance", "Même logique. Un flux, un résultat mesurable.")
    d.write_bullets(
        slide,
        d.MARGIN_L,
        d.Inches(2.35),
        d.CONTENT_W,
        d.Inches(3.2),
        [
            "Intake sinistre et routage par type, montant ou score fraude",
            "Règlement ou décaissement sinistre : blocage au-dessus du seuil jusqu'à validation",
            "Changements de police en production : prime, garanties, bénéficiaires",
            "Reporting régulateur ou partenaire : collecte, validation, transmission tracée",
        ],
        size=13,
    )
    d.add_notes(
        slide,
        "Transition : la plupart savent qu'il faut automatiser ; moins planifient la responsabilité à vitesse machine.",
    )


def slide_impact(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Impact métier", "Quand l'automatisation a un propriétaire clair")
    rows = [
        ("Réponse client plus rapide", "Onboarding, sinistres, paiements"),
        ("Coût par transaction plus bas", "Moins de reprise manuelle et de relances"),
        ("Moins d'erreurs opérationnelles", "Les règles n'oublient pas les étapes un vendredi"),
        ("Capacité de croissance", "La même équipe traite plus sans s'épuiser"),
    ]
    y = 2.35
    for head, body in rows:
        d.add_round_card(slide, d.MARGIN_L, d.Inches(y), d.CONTENT_W, d.Inches(0.82))
        d.write_text(slide, d.MARGIN_L + d.Inches(0.2), d.Inches(y + 0.14), d.Inches(3.2), d.Inches(0.3), head, size=12, bold=True, color=d.TEAL)
        d.write_text(slide, d.MARGIN_L + d.Inches(3.5), d.Inches(y + 0.16), d.Inches(5.0), d.Inches(0.45), body, size=12, color=d.MUTED)
        y += 0.95
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(6.15),
        d.CONTENT_W,
        d.Inches(0.45),
        "Attention : l'automatisation pensée seulement pour la vitesse, sans contrôle, peut augmenter le risque.",
        size=11,
        color=d.DIM,
    )
    d.add_notes(slide, "Mise en garde honnête qui introduit l'Acte 2.")


def slide_start_smart(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Commencer intelligemment", "Pas besoin d'une transformation sur cinq ans")
    steps = [
        ("1", "Un flux critique", "Paiement, sinistre, KYC ou équivalent"),
        ("2", "Pilote 4 à 8 semaines", "Règles, intégration, approbations, succès mesurable"),
        ("3", "Prouver le ROI", "Puis étendre au flux suivant"),
    ]
    y = 2.45
    for num, head, body in steps:
        d.add_round_card(slide, d.MARGIN_L, d.Inches(y), d.CONTENT_W, d.Inches(0.95))
        d.write_text(slide, d.MARGIN_L + d.Inches(0.2), d.Inches(y + 0.18), d.Inches(0.35), d.Inches(0.3), num, size=16, bold=True, color=d.TEAL)
        d.write_text(slide, d.MARGIN_L + d.Inches(0.6), d.Inches(y + 0.15), d.Inches(3.5), d.Inches(0.3), head, size=13, bold=True)
        d.write_text(slide, d.MARGIN_L + d.Inches(4.2), d.Inches(y + 0.18), d.Inches(4.2), d.Inches(0.55), body, size=12, color=d.MUTED)
        y += 1.08
    d.add_notes(
        slide,
        "Clôture Acte 1 : l'automatisation vaut le coup. La question est comment la faire accepter par risque et audit.",
    )


def slide_when_wrong(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Acte 2 · Le risque", "Quand l'automatisation dérape")
    d.write_bullets(
        slide,
        d.MARGIN_L,
        d.Inches(2.35),
        d.CONTENT_W,
        d.Inches(3.0),
        [
            "Un mauvais virement s'exécute instantanément. Aucune pause naturelle pour le corriger.",
            "Sinistre réglé sans les contrôles exigés par votre politique",
            "Changement de limite appliqué avant qu'une personne habilitée l'ait vu",
            "Qui a autorisé cela mardi à 14 h ? Logs éparpillés, tickets, chat.",
        ],
        size=13,
    )
    d.write_text(slide, d.MARGIN_L, d.Inches(5.65), d.CONTENT_W, d.Inches(0.45), "La vitesse sans preuve est une responsabilité.", size=16, bold=True, color=d.TEAL)
    d.add_notes(
        slide,
        "Pourquoi beaucoup de projets d'automatisation bloquent : le risque refuse de signer sans preuve.",
    )


def slide_accountability(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Le déficit de responsabilité", "Trois coûts cachés")
    cards = [
        ("Charge audit", "Semaines à reconstruire les événements à partir de dossiers incomplets"),
        ("Paralysie comité", "Le comité risque bloque le projet suivant. Aucune trace du précédent."),
        ("Exposition réglementaire", "Quand argent et données clients bougent sans preuve"),
    ]
    w = d.Inches(2.75)
    y = d.Inches(2.45)
    h = d.Inches(2.55)
    for i, (head, body) in enumerate(cards):
        x = d.MARGIN_L + i * (w + d.Inches(0.22))
        d.add_round_card(slide, x, y, w, h)
        d.write_text(slide, x + d.Inches(0.18), y + d.Inches(0.22), w - d.Inches(0.3), d.Inches(0.45), head, size=13, bold=True, color=d.TEAL)
        d.write_text(slide, x + d.Inches(0.18), y + d.Inches(0.75), w - d.Inches(0.3), d.Inches(1.5), body, size=11, color=d.MUTED)
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(5.35),
        d.CONTENT_W,
        d.Inches(0.9),
        "Les logs disent qu'il s'est passé quelque chose. Ils répondent rarement : quelle règle, quel approbateur, "
        "et si l'enregistrement est vérifiable plus tard.",
        size=12,
        color=d.MUTED,
    )
    d.add_notes(slide, "Écart qu'Aegis comble.")


def slide_systems_ok(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(
        slide,
        "Vos systèmes fonctionnent",
        "Ce n'est pas une panne. C'est la condition pour le prochain projet d'automatisation.",
    )
    d.write_bullets(
        slide,
        d.MARGIN_L,
        d.Inches(2.35),
        d.CONTENT_W,
        d.Inches(3.5),
        [
            "Si le core tourne bien aujourd'hui, tant mieux. La question est ce que vous automatisez demain",
            "Beaucoup veulent IA, scripts ou orchestration. Le risque bloque car personne ne prouve qui a autorisé quoi",
            "Sans registre signé des approbations, l'audit se reconstruit en semaines, pas en heures",
            "L'automatisation gouvernée permet de dire oui au projet suivant parce que la preuve existe dès le jour 1",
        ],
        size=13,
    )
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(5.85),
        d.CONTENT_W,
        d.Inches(0.55),
        "Des systèmes stables sont un atout. Salanor vous aide à ajouter l'automatisation sans parier l'institution.",
        size=12,
        color=d.TEAL,
        bold=True,
    )
    d.add_notes(slide, "Réponse à : nos systèmes n'ont aucun problème. Ne dites jamais que leur stack est cassée.")


def slide_regulatory(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(
        slide,
        "Réglementation et confiance",
        "La responsabilité s'applique, que vous automatisiez ou non",
    )
    d.write_bullets(
        slide,
        d.MARGIN_L,
        d.Inches(2.35),
        d.Inches(4.3),
        d.Inches(3.5),
        [
            "POPIA (Afrique du Sud, si pertinent) : traitement licite et responsabilité",
            "LBC-FT / KYC : qui a approuvé les exceptions",
            "Audit interne et contrôles type SOC : preuve que les contrôles ont fonctionné",
            "EU AI Act, SOC 2 : pertinents pour les groupes internationaux",
        ],
        size=12,
    )
    d.add_round_card(slide, d.Inches(5.15), d.Inches(2.35), d.Inches(3.95), d.Inches(2.35))
    d.write_text(
        slide,
        d.Inches(5.35),
        d.Inches(2.55),
        d.Inches(3.55),
        d.Inches(1.8),
        "Régulateurs et conseils ne demandent pas :\n« Le robot a-t-il tourné ? »\n\nIls demandent :\n« Qui était responsable,\net pouvez-vous le prouver ? »",
        size=13,
        color=d.TEAL,
        bold=True,
    )
    d.add_notes(slide, "POPIA si audience Afrique du Sud. Point universel : la preuve, pas seulement l'efficacité.")


def slide_both(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Il faut les deux", "Opérations rapides et contrôle prouvable")
    cols = [
        ("Manuel seul", "Lent, incohérent, difficile à scaler"),
        ("Automatisation sans gouvernance", "Rapide, mais aveugle et risqué"),
        ("Modèle Salanor", "Implémenter le flux + Aegis : rapide et prouvable"),
    ]
    w = d.Inches(2.75)
    y = d.Inches(2.55)
    h = d.Inches(2.35)
    for i, (head, body) in enumerate(cols):
        x = d.MARGIN_L + i * (w + d.Inches(0.22))
        d.add_round_card(slide, x, y, w, h)
        accent = d.TEAL if i == 2 else d.MUTED
        d.write_text(slide, x + d.Inches(0.18), y + d.Inches(0.22), w - d.Inches(0.3), d.Inches(0.55), head, size=12, bold=True, color=accent)
        d.write_text(slide, x + d.Inches(0.18), y + d.Inches(0.85), w - d.Inches(0.3), d.Inches(1.2), body, size=11, color=d.MUTED)
    d.add_notes(slide, "Acte 3 commence. Salanor implémente et intègre Aegis dès le premier jour.")


def slide_salanor(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Acte 3 · Salanor", "Implémentation avec gouvernance intégrée")
    d.write_bullets(
        slide,
        d.MARGIN_L,
        d.Inches(2.35),
        d.CONTENT_W,
        d.Inches(2.5),
        [
            "Concevoir le flux prioritaire avec ops, risque et IT",
            "Implémenter l'automatisation sur votre stack existante, sans remplacer le core",
            "Intégrer Aegis dès le jour 1 : règles, approbations, traces signées, exports audit",
        ],
        size=13,
    )
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(5.15),
        d.CONTENT_W,
        d.Inches(0.55),
        "Nous automatisons le flux. Aegis est intégré pour que risque et audit puissent y faire confiance.",
        size=14,
        bold=True,
        color=d.TEAL,
    )
    d.add_notes(slide, "Pas seulement éditeur logiciel : conception + implémentation + Aegis.")


def slide_aegis(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(
        slide,
        "Qu'est-ce qu'Aegis",
        "Contrôle et preuve avant qu'une action sensible s'exécute",
    )
    d.write_bullets(
        slide,
        d.MARGIN_L,
        d.Inches(2.35),
        d.Inches(4.5),
        d.Inches(2.8),
        [
            "La règle évalue : autoriser, refuser ou exiger approbation",
            "Un humain identifié approuve ou refuse si requis",
            "L'action ne part qu'après validation du contrôle",
            "Chaque étape dans un registre signé, append-only",
        ],
        size=12,
    )
    d.add_round_card(slide, d.Inches(5.15), d.Inches(2.35), d.Inches(3.95), d.Inches(2.8))
    d.write_text(slide, d.Inches(5.35), d.Inches(2.55), d.Inches(3.5), d.Inches(0.35), "Aegis n'est pas", size=12, bold=True, color=d.TEAL)
    d.write_bullets(
        slide,
        d.Inches(5.35),
        d.Inches(2.95),
        d.Inches(3.5),
        d.Inches(2.0),
        ["Un core banking ou SI sinistres", "Un SIEM ou agrégateur de logs", "Un remplacement de la conformité"],
        size=11,
        color=d.MUTED,
    )
    d.add_notes(slide, "BYOK en production ; Workflow Bridge pour orchestrateurs en démo.")


def slide_client_limits(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(
        slide,
        "Limites et règles par client",
        "Chemin honnête : règles par montant aujourd'hui, votre modèle en déploiement",
    )
    phases = [
        (
            "Aujourd'hui (livré)",
            [
                "Seuils par montant et type d'action",
                "Contexte client complet pour l'approbateur : compte, bénéficiaire, segment",
                "Plusieurs politiques actives en parallèle",
            ],
        ),
        (
            "Phase 2 (avec vous)",
            [
                "Règles par segment client, type de compte, listes bénéficiaires",
                "Champs métier de votre core dans la condition de règle",
            ],
        ),
        (
            "Phase 3 (optionnel)",
            [
                "Modèle de risque sur vos données pour recommander le plafond par client",
                "Le scoring reste chez vous ou se construit en cadrage avec vos équipes risque",
            ],
        ),
    ]
    w = d.Inches(2.75)
    y = d.Inches(2.35)
    h = d.Inches(3.15)
    for i, (head, bullets) in enumerate(phases):
        x = d.MARGIN_L + i * (w + d.Inches(0.22))
        d.add_round_card(slide, x, y, w, h)
        d.write_text(slide, x + d.Inches(0.15), y + d.Inches(0.15), w - d.Inches(0.25), d.Inches(0.45), head, size=11, bold=True, color=d.TEAL)
        d.write_bullets(slide, x + d.Inches(0.15), y + d.Inches(0.62), w - d.Inches(0.25), d.Inches(2.35), bullets, size=10, color=d.MUTED)
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(5.75),
        d.CONTENT_W,
        d.Inches(0.75),
        "Nous fournissons exécution et preuve. Les limites par client issues de votre modèle de risque "
        "se cartographient en déploiement avec vos équipes paiement et risque.",
        size=11,
        color=d.MUTED,
    )


def slide_four_steps(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Comment ça marche", "Quatre étapes, une chaîne de preuve")
    steps = [
        ("01", "Raccorder", "n8n, SDK, ou format ouvert APS-1"),
        ("02", "Appliquer les règles", "Autoriser, bloquer, ou mettre en attente d'approbation"),
        ("03", "Enregistrer et signer", "Registre append-only, lots de témoin"),
        ("04", "Rejouer et exporter", "Console, vérification, bundles audit"),
    ]
    w = d.Inches(2.05)
    y = d.Inches(2.55)
    h = d.Inches(2.35)
    for i, (num, head, body) in enumerate(steps):
        x = d.MARGIN_L + i * (w + d.Inches(0.18))
        d.add_round_card(slide, x, y, w, h)
        d.write_text(slide, x + d.Inches(0.12), y + d.Inches(0.15), d.Inches(0.4), d.Inches(0.25), num, size=11, bold=True, color=d.TEAL)
        d.write_text(slide, x + d.Inches(0.12), y + d.Inches(0.45), w - d.Inches(0.2), d.Inches(0.35), head, size=12, bold=True)
        d.write_text(slide, x + d.Inches(0.12), y + d.Inches(0.85), w - d.Inches(0.2), d.Inches(1.3), body, size=10, color=d.MUTED)
    d.add_notes(slide, "Pas de refonte applicative.")


def slide_demo(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Démonstration live", "Scénario banque : virement sortant au-dessus du plafond")
    d.add_round_card(slide, d.MARGIN_L, d.Inches(2.35), d.CONTENT_W, d.Inches(3.35))
    d.write_text(
        slide,
        d.MARGIN_L + d.Inches(0.25),
        d.Inches(2.55),
        d.Inches(8.0),
        d.Inches(0.45),
        "Cas : le workflow tente un virement de 2 500 USD. Règle : au-delà de 1 000 USD, approbation humaine requise.",
        size=13,
        bold=True,
    )
    d.write_bullets(
        slide,
        d.MARGIN_L + d.Inches(0.25),
        d.Inches(3.05),
        d.Inches(8.0),
        d.Inches(2.4),
        [
            "Sans gouvernance : le virement part dès que le workflow s'exécute",
            "Avec Aegis : le flux s'arrête ; l'approbateur voit montant, bénéficiaire, contexte",
            "Sans approbation : pas de virement. Chaque étape enregistrée et signée.",
            "Même mécanique pour sinistres, KYC, changements de limite.",
        ],
        size=12,
    )
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(6.0),
        d.CONTENT_W,
        d.Inches(0.35),
        "Démo : app.salanor.com · Préparer l'historique Approbations avant la réunion",
        size=11,
        color=d.DIM,
    )
    d.add_notes(slide, "Démo 10 min. Voir NOTES.md pour le script complet. Loom en secours si réseau défaillant.")


def slide_approval(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Approbation humaine", "« Proposé par le système » vs « autorisé par une personne »")
    d.write_bullets(
        slide,
        d.MARGIN_L,
        d.Inches(2.35),
        d.CONTENT_W,
        d.Inches(3.2),
        [
            "Alertes par e-mail, Slack, PagerDuty ou SMS",
            "Approbateur identifié, horodaté, trace unique",
            "Refus ou expiration : l'action ne s'exécute pas",
            "Historique complet : approuvé, refusé, expiré, avec la règle appliquée",
        ],
        size=13,
    )
    d.add_notes(slide, "Ce que les comités risque et l'audit interne attendent.")


def slide_audit(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Audit et conformité", "Des preuves pour vos revues")
    d.write_bullets(
        slide,
        d.MARGIN_L,
        d.Inches(2.35),
        d.CONTENT_W,
        d.Inches(3.2),
        [
            "Exports par période avec hash d'intégrité vérifiable",
            "Correspondance contrôles : SOC 2 et EU AI Act dans les exports (aide doc., pas certification)",
            "Relecture pas à pas dans la console",
            "Journal admin : politiques, clés, connexions, exports",
        ],
        size=12,
    )
    d.add_notes(slide, "POPIA : la responsabilité est plus simple avec un dossier complet vérifiable qu'avec des e-mails.")


def slide_offer(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(slide, "Ce que nous livrons", "4 à 8 semaines · un flux · automatisation + Aegis")
    rows = [
        ("Périmètre", "Un processus que nous automatisons pour vous : paiement, sinistre, KYC ou équivalent"),
        ("Livré", "Flux sur votre stack (n8n, API, scripts) avec règles Aegis, approbations, trace signée"),
        ("Plateforme", "Formule Aegis Team à partir de 299 USD / mois"),
        ("Implémentation", "Forfait fixe selon complexité : orchestrateur seul vs API core + BYOK"),
        ("Succès", "Le flux va plus vite. Votre audit valide la preuve sans Salanor dans la salle."),
    ]
    y = 2.35
    for head, body in rows:
        d.add_round_card(slide, d.MARGIN_L, d.Inches(y), d.CONTENT_W, d.Inches(0.72))
        d.write_text(slide, d.MARGIN_L + d.Inches(0.2), d.Inches(y + 0.12), d.Inches(1.55), d.Inches(0.28), head, size=11, bold=True, color=d.TEAL)
        d.write_text(slide, d.MARGIN_L + d.Inches(1.85), d.Inches(y + 0.14), d.Inches(6.5), d.Inches(0.45), body, size=12, color=d.MUTED)
        y += 0.82
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(6.35),
        d.CONTENT_W,
        d.Inches(0.3),
        "api.salanor.com · app.salanor.com · SDK n8n, TypeScript, Python, Go",
        size=10,
        color=d.DIM,
    )
    d.add_notes(slide, "Vous vendez implémentation + Aegis. Pas SaaS seul. Pas partenaire de conception.")


def slide_next_step(prs) -> None:
    slide = d.blank(prs)
    d.slide_header(
        slide,
        "Prochaine étape",
        "Une question, un cas concret, ou envie d'en voir plus ? Prenons rendez-vous.",
    )
    options = [
        ("15 minutes", "Intro courte, Q&R, aperçu console ou démo live"),
        ("45 minutes", "Cadrage avec ops, risque ou IT : un flux et un critère de succès"),
        ("Sur demande", "Exemple d'export audit, fiche une page, ou démo sur votre scénario"),
    ]
    w = d.Inches(2.75)
    y = d.Inches(2.45)
    h = d.Inches(2.35)
    for i, (head, body) in enumerate(options):
        x = d.MARGIN_L + i * (w + d.Inches(0.22))
        d.add_round_card(slide, x, y, w, h)
        d.write_text(slide, x + d.Inches(0.18), y + d.Inches(0.22), w - d.Inches(0.3), d.Inches(0.45), head, size=13, bold=True, color=d.TEAL)
        d.write_text(slide, x + d.Inches(0.18), y + d.Inches(0.72), w - d.Inches(0.3), d.Inches(1.45), body, size=11, color=d.MUTED)
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(5.15),
        d.CONTENT_W,
        d.Inches(0.55),
        "Contactez-moi directement ou via M. Yassine. Nous fixons un créneau et repartons avec un périmètre clair.",
        size=12,
        color=d.MUTED,
    )
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(5.85),
        d.CONTENT_W,
        d.Inches(0.45),
        "contact@salanor.com · partners@salanor.com · app.salanor.com",
        size=13,
        color=d.WHITE,
        bold=True,
    )
    d.add_notes(
        slide,
        "CTA du deck. En présentiel : fixer le cadrage 45 min avant de quitter la salle.",
    )


def slide_contact(prs) -> None:
    slide = d.blank(prs)
    d.add_rect(slide, d.Inches(0), d.Inches(7.44), d.Inches(10), d.Inches(0.06), d.TEAL_DIM)
    d.write_text(slide, d.MARGIN_L, d.Inches(0.55), d.CONTENT_W, d.Inches(0.28), "CONTACT", size=9, color=d.TEAL, bold=True)
    d.write_text(slide, d.MARGIN_L, d.Inches(1.05), d.CONTENT_W, d.Inches(0.55), "Salanor Ltd", size=28, bold=True)
    lines = [
        "Landry Bougang Fotso · Fondateur · Kigali, Rwanda",
        "",
        "www.salanor.com",
        "www.salanor.com/products/aegis",
        "app.salanor.com",
        "",
        "contact@salanor.com · partners@salanor.com",
    ]
    y = 1.85
    for line in lines:
        if line:
            d.write_text(
                slide,
                d.MARGIN_L,
                d.Inches(y),
                d.CONTENT_W,
                d.Inches(0.35),
                line,
                size=13,
                color=d.MUTED if "@" in line or "www" in line else d.WHITE,
            )
        y += 0.42 if line else 0.2
    d.write_text(
        slide,
        d.MARGIN_L,
        d.Inches(6.5),
        d.CONTENT_W,
        d.Inches(0.3),
        "Automatisation et gouvernance · Finance et assurance · Confidentiel",
        size=9,
        color=d.DIM,
    )
    d.add_notes(slide, "Remerciez M. Yassine. Proposez un appel de découverte ou une démo technique.")


def build() -> Path:
    prs = d.Presentation()
    prs.slide_width = d.Inches(10)
    prs.slide_height = d.Inches(7.5)

    slide_title(prs)
    slide_reality(prs)
    slide_benefits(prs)
    slide_banking(prs)
    slide_insurance(prs)
    slide_impact(prs)
    slide_start_smart(prs)
    slide_when_wrong(prs)
    slide_accountability(prs)
    slide_systems_ok(prs)
    slide_regulatory(prs)
    slide_both(prs)
    slide_salanor(prs)
    slide_aegis(prs)
    slide_client_limits(prs)
    slide_four_steps(prs)
    slide_demo(prs)
    slide_approval(prs)
    slide_audit(prs)
    slide_offer(prs)
    slide_next_step(prs)
    slide_contact(prs)

    OUT_REPO.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_REPO))

    OUT_YASSINE.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT_REPO, OUT_YASSINE)

    return OUT_REPO


if __name__ == "__main__":
    path = build()
    from pptx import Presentation as PptxPresentation

    prs = PptxPresentation(str(path))
    print(f"Wrote {path}")
    print(f"Slides: {len(prs.slides)}")
    print(f"Copied to {OUT_YASSINE}")
