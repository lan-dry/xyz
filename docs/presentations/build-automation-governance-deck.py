"""Generate Salanor Automation & Governance deck (.pptx).

Run: python docs/presentations/build-automation-governance-deck.py
"""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
LOGO = ROOT / "apps" / "web-marketing" / "public" / "salanor-logo.png"
OUT = Path(__file__).resolve().parent / "Salanor-Automation-Governance-Finance.pptx"
OUT_SA = Path("e:/salanor/Presentation/South-Africa/Salanor-Automation-Governance-Finance.pptx")
OUT_YASSINE_EN = Path("e:/salanor/Presentation/Yassine/Salanor-Automation-Governance-Finance.pptx")

BG = RGBColor(0x0A, 0x0C, 0x0B)
SURFACE = RGBColor(0x14, 0x19, 0x17)
BORDER = RGBColor(0x27, 0x3A, 0x38)
TEAL = RGBColor(0x43, 0xC2, 0xA9)
TEAL_DIM = RGBColor(0x0D, 0x35, 0x35)
WHITE = RGBColor(0xE8, 0xED, 0xE9)
MUTED = RGBColor(0x9E, 0xB0, 0xAC)
DIM = RGBColor(0x71, 0x85, 0x82)

FONT = "Segoe UI"
MARGIN_L = Inches(0.65)
CONTENT_W = Inches(8.7)


def set_slide_bg(slide, color: RGBColor = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill: RGBColor, line: RGBColor | None = None) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()


def add_round_card(slide, left, top, width, height) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = SURFACE
    shape.line.color.rgb = BORDER
    shape.line.width = Pt(0.75)
    return shape


def write_text(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: int = 14,
    bold: bool = False,
    color: RGBColor = WHITE,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT


def write_bullets(slide, left, top, width, height, items: list[str], size: int = 12, color: RGBColor = WHITE) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = FONT


def slide_header(slide, eyebrow: str, title: str, subtitle: str | None = None) -> None:
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), MARGIN_L, Inches(0.42), width=Inches(0.38))
    write_text(slide, Inches(1.08), Inches(0.48), Inches(1.2), Inches(0.25), "SALANOR", size=9, color=DIM, bold=True)
    write_text(slide, MARGIN_L, Inches(0.82), CONTENT_W, Inches(0.28), eyebrow.upper(), size=9, color=TEAL, bold=True)
    write_text(slide, MARGIN_L, Inches(1.05), CONTENT_W, Inches(0.85), title, size=24, bold=True)
    add_rect(slide, MARGIN_L, Inches(1.88), Inches(1.1), Inches(0.045), TEAL)
    if subtitle:
        write_text(slide, MARGIN_L, Inches(2.02), CONTENT_W, Inches(0.55), subtitle, size=13, color=MUTED)


def slide_footer(slide, text: str) -> None:
    write_text(slide, MARGIN_L, Inches(7.05), CONTENT_W, Inches(0.3), text, size=9, color=DIM)


def add_notes(slide, text: str) -> None:
    notes = slide.notes_slide.notes_text_frame
    notes.text = text


def blank(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    return slide


def slide_title(prs: Presentation) -> None:
    slide = blank(prs)
    add_rect(slide, Inches(0), Inches(0), Inches(10), Inches(0.06), TEAL_DIM)
    write_text(slide, MARGIN_L, Inches(0.55), CONTENT_W, Inches(0.28), "SALANOR · CONFIDENTIAL · 2026", size=9, color=TEAL, bold=True)
    write_text(
        slide,
        MARGIN_L,
        Inches(1.15),
        Inches(8.5),
        Inches(1.4),
        "Automation & governance\nfor finance & insurance",
        size=36,
        bold=True,
    )
    write_text(
        slide,
        MARGIN_L,
        Inches(2.75),
        Inches(8.2),
        Inches(1.0),
        "We automate your critical workflows and build in Aegis from day one: speed, control, and proof for payments, claims, and KYC.",
        size=15,
        color=MUTED,
    )
    write_text(slide, MARGIN_L, Inches(6.35), CONTENT_W, Inches(0.35), "Landry Bougang Fotso · Founder, Salanor Ltd · Kigali", size=11, color=DIM)
    add_notes(slide, "Introduce yourself. Three-part story: why automate, the gap, Salanor + Aegis. Finance & insurance only.")


def slide_reality(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "Act 1 · The opportunity", "Many institutions still run critical work manually")
    write_bullets(
        slide,
        MARGIN_L,
        Inches(2.35),
        CONTENT_W,
        Inches(3.5),
        [
            "Email chains, spreadsheets, paper handoffs, duplicate data entry",
            "Slow turnaround: clients wait, staff chase approvals",
            "Inconsistent decisions depending on who is on shift",
            "Errors that hurt more as volume grows",
            "Process knowledge walks out the door when people leave",
        ],
        size=13,
    )
    write_text(
        slide,
        MARGIN_L,
        Inches(5.95),
        CONTENT_W,
        Inches(0.55),
        "Digital expectations are rising globally. Manual back-office processes become the bottleneck.",
        size=12,
        color=MUTED,
    )
    slide_footer(slide, "Automation & governance · Finance & insurance")
    add_notes(slide, "Not Africa-only — many regulated institutions worldwide still manual. Transition: one workflow done properly is enough to start.")


def slide_benefits(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(
        slide,
        "What automation brings",
        "Structured workflow automation — not “replace everyone with AI”",
    )
    cards = [
        ("Speed", "Routine steps in minutes, not hours or days"),
        ("Consistency", "Same rule, same path, every time"),
        ("Scale", "More volume without linear headcount growth"),
        ("Focus", "Staff on judgment, not copy-paste"),
    ]
    w = Inches(2.05)
    gap = Inches(0.18)
    y = Inches(2.45)
    h = Inches(2.2)
    for i, (head, body) in enumerate(cards):
        x = MARGIN_L + i * (w + gap)
        add_round_card(slide, x, y, w, h)
        write_text(slide, x + Inches(0.15), y + Inches(0.18), w - Inches(0.25), Inches(0.35), head, size=13, bold=True, color=TEAL)
        write_text(slide, x + Inches(0.15), y + Inches(0.58), w - Inches(0.25), Inches(1.4), body, size=11, color=MUTED)
    write_text(
        slide,
        MARGIN_L,
        Inches(5.0),
        CONTENT_W,
        Inches(0.5),
        "Automation removes friction around sensitive decisions. It does not remove humans from them.",
        size=12,
        color=MUTED,
    )
    add_notes(slide, "Define automation as rules + routing + notifications. Avoid AI hype.")


def slide_banking(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "Banking", "High-value starting points. Pick one workflow.")
    write_bullets(
        slide,
        MARGIN_L,
        Inches(2.35),
        CONTENT_W,
        Inches(3.2),
        [
            "Outbound payments above threshold: approval before execution",
            "KYC / AML document checks: route exceptions, flag incomplete files",
            "Limit or account status changes: dual control before core update",
            "Reconciliation alerts: surface mismatches with context to the right team",
        ],
        size=13,
    )
    add_notes(slide, "Pick the workflow that keeps ops or compliance awake at night.")


def slide_insurance(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "Insurance", "Same logic. One workflow, measurable outcome.")
    write_bullets(
        slide,
        MARGIN_L,
        Inches(2.35),
        CONTENT_W,
        Inches(3.2),
        [
            "Claims intake and routing by type, amount, or fraud score",
            "Claim settlement or disbursement: block above threshold until sign-off",
            "Policy changes in production: premium, coverage, beneficiary updates",
            "Regulatory or partner reporting: gather, validate, transmit with record",
        ],
        size=13,
    )
    add_notes(slide, "Bridge: most know they should automate; fewer plan for accountability at machine speed.")


def slide_impact(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "Business impact", "When automation has clear ownership")
    rows = [
        ("Faster customer response", "Onboarding, claims, payments"),
        ("Lower cost per transaction", "Less manual rework and chasing"),
        ("Fewer operational errors", "Rules do not forget steps on a Friday"),
        ("Room to grow", "Same team handles more without burning out"),
    ]
    y = 2.35
    for head, body in rows:
        add_round_card(slide, MARGIN_L, Inches(y), CONTENT_W, Inches(0.82))
        write_text(slide, MARGIN_L + Inches(0.2), Inches(y + 0.14), Inches(3.2), Inches(0.3), head, size=12, bold=True, color=TEAL)
        write_text(slide, MARGIN_L + Inches(3.5), Inches(y + 0.16), Inches(5.0), Inches(0.45), body, size=12, color=MUTED)
        y += 0.95
    write_text(
        slide,
        MARGIN_L,
        Inches(6.15),
        CONTENT_W,
        Inches(0.45),
        "Caveat: automation built only for speed, without control, can increase risk.",
        size=11,
        color=DIM,
    )
    add_notes(slide, "Honest caveat leads into Act 2.")


def slide_start_smart(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "Start smart", "No five-year transformation required")
    steps = [
        ("1", "One critical workflow", "Payment, claim, KYC, or similar"),
        ("2", "4–8 week pilot", "Rules, integration, approvals, measurable success"),
        ("3", "Prove ROI", "Then expand to the next workflow"),
    ]
    y = 2.45
    for num, head, body in steps:
        add_round_card(slide, MARGIN_L, Inches(y), CONTENT_W, Inches(0.95))
        write_text(slide, MARGIN_L + Inches(0.2), Inches(y + 0.18), Inches(0.35), Inches(0.3), num, size=16, bold=True, color=TEAL)
        write_text(slide, MARGIN_L + Inches(0.6), Inches(y + 0.15), Inches(3.5), Inches(0.3), head, size=13, bold=True)
        write_text(slide, MARGIN_L + Inches(4.2), Inches(y + 0.18), Inches(4.2), Inches(0.55), body, size=12, color=MUTED)
        y += 1.08
    add_notes(slide, "Act 1 close: automation is worth doing — the question is how to do it so risk and audit say yes.")


def slide_when_wrong(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "Act 2 · The gap", "When automation goes wrong")
    write_bullets(
        slide,
        MARGIN_L,
        Inches(2.35),
        CONTENT_W,
        Inches(3.0),
        [
            "Wrong payment executes instantly. No natural pause to catch it.",
            "Claim settled without the checks your policy requires",
            "Limit change applies before anyone with authority has seen it",
            "Who authorized this Tuesday at 2 p.m.? Scattered logs, tickets, chat.",
        ],
        size=13,
    )
    write_text(slide, MARGIN_L, Inches(5.65), CONTENT_W, Inches(0.45), "Speed without proof is liability.", size=16, bold=True, color=TEAL)
    add_notes(slide, "Why many automation projects stall — risk blocks because nobody wants to sign without evidence.")


def slide_accountability(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "The accountability gap", "Three hidden costs")
    cards = [
        ("Audit drag", "Weeks reconstructing events from incomplete records"),
        ("Committee paralysis", "Risk boards block the next project. No trail from the last one."),
        ("Regulatory exposure", "When money and client data move automatically without proof"),
    ]
    w = Inches(2.75)
    y = Inches(2.45)
    h = Inches(2.55)
    for i, (head, body) in enumerate(cards):
        x = MARGIN_L + i * (w + Inches(0.22))
        add_round_card(slide, x, y, w, h)
        write_text(slide, x + Inches(0.18), y + Inches(0.22), w - Inches(0.3), Inches(0.45), head, size=13, bold=True, color=TEAL)
        write_text(slide, x + Inches(0.18), y + Inches(0.75), w - Inches(0.3), Inches(1.5), body, size=11, color=MUTED)
    write_text(
        slide,
        MARGIN_L,
        Inches(5.35),
        CONTENT_W,
        Inches(0.9),
        "Logs say something happened. They rarely answer: which rule applied, which person approved, and whether the record can be verified later.",
        size=12,
        color=MUTED,
    )
    add_notes(slide, "Gap Aegis closes.")


def slide_systems_ok(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(
        slide,
        "Your systems work",
        "This is not a fix for outages. It unlocks the next automation project.",
    )
    write_bullets(
        slide,
        MARGIN_L,
        Inches(2.35),
        CONTENT_W,
        Inches(3.5),
        [
            "If core systems run well today, good. The question is what you automate next",
            "Many teams want AI, scripts, or orchestration. Risk blocks because nobody can prove who authorized what",
            "Without a signed approval register, audit reconstruction takes weeks, not hours",
            "Governed automation lets you say yes to the next project because proof exists from day one",
        ],
        size=13,
    )
    write_text(
        slide,
        MARGIN_L,
        Inches(5.85),
        CONTENT_W,
        Inches(0.55),
        "Stable systems are an advantage. Salanor helps you add automation without betting the institution.",
        size=12,
        color=TEAL,
        bold=True,
    )
    add_notes(slide, "Answer to: our systems have no issues. Never say their stack is broken.")


def slide_client_limits(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(
        slide,
        "Limits and rules per client",
        "Honest path: amount rules today, your data model in deployment",
    )
    phases = [
        ("Today (shipped)", [
            "Thresholds by amount and action type",
            "Full client context for the approver: account, beneficiary, segment",
            "Multiple active policies in parallel",
        ]),
        ("Phase 2 (with you)", [
            "Rules by client segment, account type, beneficiary lists",
            "Your core banking fields inside the rule condition",
        ]),
        ("Phase 3 (optional)", [
            "Risk model on your data to recommend limit per client",
            "Scoring stays with you or is built with your risk team in scoping",
        ]),
    ]
    w = Inches(2.75)
    y = Inches(2.35)
    h = Inches(3.15)
    for i, (head, bullets) in enumerate(phases):
        x = MARGIN_L + i * (w + Inches(0.22))
        add_round_card(slide, x, y, w, h)
        write_text(slide, x + Inches(0.15), y + Inches(0.15), w - Inches(0.25), Inches(0.45), head, size=11, bold=True, color=TEAL)
        write_bullets(slide, x + Inches(0.15), y + Inches(0.62), w - Inches(0.25), Inches(2.35), bullets, size=10, color=MUTED)
    write_text(
        slide,
        MARGIN_L,
        Inches(5.75),
        CONTENT_W,
        Inches(0.75),
        "We provide execution and proof. Per-client limits from your risk model are mapped during deployment with payment and risk teams.",
        size=11,
        color=MUTED,
    )


def slide_regulatory(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(
        slide,
        "Regulatory & trust",
        "Accountability applies whether you automate or not",
    )
    write_bullets(
        slide,
        MARGIN_L,
        Inches(2.35),
        Inches(4.3),
        Inches(3.5),
        [
            "POPIA (South Africa): lawful processing and accountability",
            "AML / KYC: who approved exceptions",
            "Internal audit and SOC-style controls: evidence controls operated",
            "EU AI Act, SOC 2: relevant for international groups",
        ],
        size=12,
    )
    add_round_card(slide, Inches(5.15), Inches(2.35), Inches(3.95), Inches(2.35))
    write_text(
        slide,
        Inches(5.35),
        Inches(2.55),
        Inches(3.55),
        Inches(1.8),
        "Regulators and boards do not ask:\n“Did the robot run?”\n\nThey ask:\n“Who was responsible,\nand can you prove it?”",
        size=13,
        color=TEAL,
        bold=True,
    )
    add_notes(slide, "POPIA if SA audience. Universal point: evidence not efficiency alone.")


def slide_both(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "You need both", "Fast operations and provable control")
    cols = [
        ("Manual only", "Slow, inconsistent, hard to scale"),
        ("Automation, no governance", "Fast, but blind and risky"),
        ("Salanor model", "Implement workflow + Aegis: fast and provable"),
    ]
    w = Inches(2.75)
    y = Inches(2.55)
    h = Inches(2.35)
    for i, (head, body) in enumerate(cols):
        x = MARGIN_L + i * (w + Inches(0.22))
        add_round_card(slide, x, y, w, h)
        accent = TEAL if i == 2 else MUTED
        write_text(slide, x + Inches(0.18), y + Inches(0.22), w - Inches(0.3), Inches(0.55), head, size=12, bold=True, color=accent)
        write_text(slide, x + Inches(0.18), y + Inches(0.85), w - Inches(0.3), Inches(1.2), body, size=11, color=MUTED)
    add_notes(slide, "Act 3 begins. Salanor implements + embeds Aegis from day one.")


def slide_salanor(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "Act 3 · Salanor", "Implementation with governance built in")
    write_bullets(
        slide,
        MARGIN_L,
        Inches(2.35),
        CONTENT_W,
        Inches(2.5),
        [
            "Design the priority workflow with ops, risk, and IT",
            "Implement automation on your existing stack — no core replacement",
            "Include Aegis from day one: rules, approvals, signed traces, audit exports",
        ],
        size=13,
    )
    write_text(
        slide,
        MARGIN_L,
        Inches(5.15),
        CONTENT_W,
        Inches(0.55),
        "We automate the workflow. Aegis is built in so risk and audit can trust it.",
        size=14,
        bold=True,
        color=TEAL,
    )
    add_notes(slide, "Not only software vendor — design + implement + Aegis.")


def slide_aegis(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(
        slide,
        "What Aegis is",
        "Control and proof before a sensitive action executes",
    )
    write_bullets(
        slide,
        MARGIN_L,
        Inches(2.35),
        Inches(4.5),
        Inches(2.8),
        [
            "Rule evaluates → allow, deny, or require approval",
            "Named human approves or refuses when required",
            "Action proceeds only after the gate clears",
            "Every step in a signed, append-only register",
        ],
        size=12,
    )
    add_round_card(slide, Inches(5.15), Inches(2.35), Inches(3.95), Inches(2.8))
    write_text(slide, Inches(5.35), Inches(2.55), Inches(3.5), Inches(0.35), "Aegis is not", size=12, bold=True, color=TEAL)
    write_bullets(
        slide,
        Inches(5.35),
        Inches(2.95),
        Inches(3.5),
        Inches(2.0),
        ["Core banking or claims platform", "SIEM or log aggregator", "Replacement for compliance"],
        size=11,
        color=MUTED,
    )
    add_notes(slide, "BYOK for production; Workflow Bridge for orchestrators in demo.")


def slide_four_steps(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "How it works", "Four steps, one chain of proof")
    steps = [
        ("01", "Connect", "n8n, SDK, or open APS-1 format"),
        ("02", "Apply rules", "Allow, block, or pause for approval"),
        ("03", "Record & sign", "Append-only ledger, witness batches"),
        ("04", "Replay & export", "Console, verification, audit bundles"),
    ]
    w = Inches(2.05)
    y = Inches(2.55)
    h = Inches(2.35)
    for i, (num, head, body) in enumerate(steps):
        x = MARGIN_L + i * (w + Inches(0.18))
        add_round_card(slide, x, y, w, h)
        write_text(slide, x + Inches(0.12), y + Inches(0.15), Inches(0.4), Inches(0.25), num, size=11, bold=True, color=TEAL)
        write_text(slide, x + Inches(0.12), y + Inches(0.45), w - Inches(0.2), Inches(0.35), head, size=12, bold=True)
        write_text(slide, x + Inches(0.12), y + Inches(0.85), w - Inches(0.2), Inches(1.3), body, size=10, color=MUTED)
    add_notes(slide, "No rip-and-replace.")


def slide_demo(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "Live demo", "Bank scenario: outbound transfer above limit")
    add_round_card(slide, MARGIN_L, Inches(2.35), CONTENT_W, Inches(3.35))
    write_text(
        slide,
        MARGIN_L + Inches(0.25),
        Inches(2.55),
        Inches(8.0),
        Inches(0.45),
        "Case: workflow attempts USD 2,500 transfer. Rule: above USD 1,000 requires human approval.",
        size=13,
        bold=True,
    )
    write_bullets(
        slide,
        MARGIN_L + Inches(0.25),
        Inches(3.05),
        Inches(8.0),
        Inches(2.4),
        [
            "Without governance → payment goes when workflow runs",
            "With Aegis → workflow stops; approver sees amount, beneficiary, context",
            "No approval → no payment. Every step recorded and signed.",
            "Same pattern for claims, KYC, limit changes.",
        ],
        size=12,
    )
    write_text(slide, MARGIN_L, Inches(6.0), CONTENT_W, Inches(0.35), "Demo: app.salanor.com · Prepare Approvals history before the meeting", size=11, color=DIM)
    add_notes(slide, "10 min demo. See NOTES.md for full script. Loom backup if network fails.")


def slide_approval(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "Human approval", "“Proposed by the system” vs “authorized by a person”")
    write_bullets(
        slide,
        MARGIN_L,
        Inches(2.35),
        CONTENT_W,
        Inches(3.2),
        [
            "Alerts via email, Slack, PagerDuty, or SMS",
            "Named approver, timestamp, one trace",
            "Refusal or timeout → action does not execute",
            "Full history: approved, refused, expired — with rule applied",
        ],
        size=13,
    )
    add_notes(slide, "What risk committees and internal audit need.")


def slide_audit(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "Audit & compliance", "Evidence for your reviews")
    write_bullets(
        slide,
        MARGIN_L,
        Inches(2.35),
        CONTENT_W,
        Inches(3.2),
        [
            "Exports by period with verifiable integrity hash",
            "Control mappings: SOC 2 & EU AI Act in exports today (documentation aid, not certification)",
            "Step-by-step replay in the console",
            "Admin audit log: policies, keys, connections, exports",
        ],
        size=12,
    )
    add_notes(slide, "POPIA: accountability easier with complete verifiable record vs email trails.")


def slide_offer(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(slide, "What we deliver", "4 to 8 weeks · one workflow · automation + Aegis")
    rows = [
        ("Scope", "One process we automate for you: payment, claim, KYC, or similar"),
        ("Built", "Workflow on your stack (n8n, API, scripts) with Aegis rules, approvals, signed trace"),
        ("Platform", "Aegis Team tier from USD 299 / month"),
        ("Implementation", "Fixed fee by complexity: orchestrator only vs core API + BYOK"),
        ("Success", "Workflow runs faster. Your audit team validates the proof without Salanor in the room."),
    ]
    y = 2.35
    for head, body in rows:
        add_round_card(slide, MARGIN_L, Inches(y), CONTENT_W, Inches(0.72))
        write_text(slide, MARGIN_L + Inches(0.2), Inches(y + 0.12), Inches(1.55), Inches(0.28), head, size=11, bold=True, color=TEAL)
        write_text(slide, MARGIN_L + Inches(1.85), Inches(y + 0.14), Inches(6.5), Inches(0.45), body, size=12, color=MUTED)
        y += 0.82
    write_text(slide, MARGIN_L, Inches(6.35), CONTENT_W, Inches(0.3), "api.salanor.com · app.salanor.com · n8n, TypeScript, Python, Go SDKs", size=10, color=DIM)
    add_notes(slide, "You sell implementation + Aegis. Not SaaS alone. Not design partner.")


def slide_next_step(prs: Presentation) -> None:
    slide = blank(prs)
    slide_header(
        slide,
        "Next step",
        "Questions, a use case, or want to see more? Book a call.",
    )
    options = [
        ("15 minutes", "Short intro, Q&A, console preview or live demo"),
        ("45 minutes", "Scoping with ops, risk, or IT: pick one workflow and define success"),
        ("On request", "Sample audit export, one-pager, or demo on your scenario"),
    ]
    w = Inches(2.75)
    y = Inches(2.45)
    h = Inches(2.35)
    for i, (head, body) in enumerate(options):
        x = MARGIN_L + i * (w + Inches(0.22))
        add_round_card(slide, x, y, w, h)
        write_text(slide, x + Inches(0.18), y + Inches(0.22), w - Inches(0.3), Inches(0.45), head, size=13, bold=True, color=TEAL)
        write_text(slide, x + Inches(0.18), y + Inches(0.72), w - Inches(0.3), Inches(1.45), body, size=11, color=MUTED)
    write_text(
        slide,
        MARGIN_L,
        Inches(5.15),
        CONTENT_W,
        Inches(0.55),
        "Email me directly. If someone introduced you, they can help set up the call.",
        size=12,
        color=MUTED,
    )
    write_text(
        slide,
        MARGIN_L,
        Inches(5.85),
        CONTENT_W,
        Inches(0.45),
        "contact@salanor.com · partners@salanor.com · app.salanor.com",
        size=13,
        color=WHITE,
        bold=True,
    )
    add_notes(slide, "Shared deck CTA. In person: book the 45 min scoping before leaving the room.")


def slide_contact(prs: Presentation) -> None:
    slide = blank(prs)
    add_rect(slide, Inches(0), Inches(7.44), Inches(10), Inches(0.06), TEAL_DIM)
    write_text(slide, MARGIN_L, Inches(0.55), CONTENT_W, Inches(0.28), "CONTACT", size=9, color=TEAL, bold=True)
    write_text(slide, MARGIN_L, Inches(1.05), CONTENT_W, Inches(0.55), "Salanor Ltd", size=28, bold=True)
    write_lines = [
        "Landry Bougang Fotso · Founder · Kigali, Rwanda",
        "",
        "www.salanor.com",
        "www.salanor.com/products/aegis",
        "app.salanor.com",
        "",
        "contact@salanor.com · partners@salanor.com",
    ]
    y = 1.85
    for line in write_lines:
        if line:
            write_text(slide, MARGIN_L, Inches(y), CONTENT_W, Inches(0.35), line, size=13, color=MUTED if "@" in line or "www" in line else WHITE)
        y += 0.42 if line else 0.2
    write_text(slide, MARGIN_L, Inches(6.5), CONTENT_W, Inches(0.3), "Automation & governance · Finance & insurance · Confidential", size=9, color=DIM)
    add_notes(slide, "Thank them. Offer discovery call or technical deep-dive.")


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    slide_title(prs)
    slide_reality(prs)
    slide_benefits(prs)
    slide_banking(prs)
    slide_insurance(prs)
    slide_impact(prs)
    slide_start_smart(prs)
    slide_when_wrong(prs)
    slide_accountability(prs)
    slide_systems_ok(prs)
    slide_regulatory(prs)
    slide_both(prs)
    slide_salanor(prs)
    slide_aegis(prs)
    slide_client_limits(prs)
    slide_four_steps(prs)
    slide_demo(prs)
    slide_approval(prs)
    slide_audit(prs)
    slide_offer(prs)
    slide_next_step(prs)
    slide_contact(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))

    for dest in (OUT_SA, OUT_YASSINE_EN):
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(OUT, dest)

    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
    print(f"Copied to {OUT_SA}")
    print(f"Copied to {OUT_YASSINE_EN}")
