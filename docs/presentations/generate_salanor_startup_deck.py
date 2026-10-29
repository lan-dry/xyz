from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("Salanor-Startup-Pitch-Deck.pptx")

W = 13.333
H = 7.5

BG = RGBColor(8, 17, 22)
SURFACE = RGBColor(15, 29, 34)
SURFACE_2 = RGBColor(20, 39, 44)
TEAL = RGBColor(67, 194, 169)
TEAL_DARK = RGBColor(13, 53, 53)
WHITE = RGBColor(241, 247, 246)
MUTED = RGBColor(158, 177, 176)
DIM = RGBColor(105, 127, 127)
LINE = RGBColor(42, 65, 68)
RED = RGBColor(230, 106, 106)
AMBER = RGBColor(231, 182, 95)

FONT = "Aptos"
FONT_DISPLAY = "Aptos Display"


def add_bg(slide):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(H)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=18,
    color=WHITE,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
    line_spacing=1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_text(
    slide,
    lines,
    x,
    y,
    w,
    h,
    *,
    size=18,
    color=WHITE,
    bullet=False,
    gap=8,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.0
        p.level = 0
        if bullet:
            p.text = f"•  {line}"
        else:
            p.text = line
        for run in p.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, *, fill=SURFACE, line=LINE, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        kind, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.5):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x1),
        Inches(y1),
        Inches(max(0.01, x2 - x1)),
        Inches(max(0.01, y2 - y1)),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def pill(slide, text, x, y, w, *, fill=TEAL_DARK, color=TEAL):
    rect(slide, x, y, w, 0.34, fill=fill, line=fill)
    add_text(
        slide,
        text.upper(),
        x,
        y + 0.01,
        w,
        0.28,
        size=8,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def brand(slide, section=None, number=None):
    add_text(slide, "SALANOR", 0.65, 0.34, 1.5, 0.25, size=10, bold=True)
    if section:
        add_text(slide, section.upper(), 2.0, 0.34, 4.5, 0.25, size=9, color=DIM)
    if number is not None:
        add_text(
            slide,
            str(number),
            12.2,
            7.05,
            0.45,
            0.2,
            size=8,
            color=DIM,
            align=PP_ALIGN.RIGHT,
        )


def footer(slide, text="Confidential | Salanor Ltd | 2026"):
    add_text(slide, text, 0.65, 7.05, 5.5, 0.18, size=7, color=DIM)


def slide_title(slide, eyebrow, title, subtitle=None, number=None):
    brand(slide, eyebrow, number)
    add_text(
        slide,
        title,
        0.65,
        0.88,
        11.9,
        0.72,
        size=28,
        bold=True,
        font=FONT_DISPLAY,
    )
    if subtitle:
        add_text(slide, subtitle, 0.65, 1.64, 11.6, 0.58, size=13, color=MUTED)
    footer(slide)


def card(
    slide,
    x,
    y,
    w,
    h,
    title,
    body,
    *,
    accent=TEAL,
    number=None,
    body_size=12,
):
    rect(slide, x, y, w, h)
    if number is not None:
        add_text(slide, number, x + 0.24, y + 0.22, 0.42, 0.3, size=9, color=accent, bold=True)
        tx = x + 0.72
    else:
        tx = x + 0.28
    add_text(slide, title, tx, y + 0.19, w - (tx - x) - 0.25, 0.54, size=14, bold=True)
    add_text(
        slide,
        body,
        x + 0.28,
        y + 0.77,
        w - 0.56,
        h - 0.98,
        size=body_size,
        color=MUTED,
        line_spacing=1.05,
    )


def arrow(slide, x, y, w=0.5, color=TEAL):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.32)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


# 1. Cover
slide = prs.slides.add_slide(blank)
add_bg(slide)
add_text(slide, "SALANOR", 0.72, 0.62, 2.4, 0.4, size=14, bold=True)
pill(slide, "Startup pitch | 2026", 10.35, 0.62, 2.15)
add_text(
    slide,
    "Control and proof\nfor enterprise AI",
    0.72,
    1.55,
    8.8,
    1.65,
    size=38,
    bold=True,
    font=FONT_DISPLAY,
)
add_text(
    slide,
    "Salanor helps regulated teams use AI without losing control of data, decisions, or actions.",
    0.76,
    3.48,
    7.6,
    0.88,
    size=17,
    color=MUTED,
    line_spacing=1.05,
)
line(slide, 0.75, 5.36, 7.2, 5.39, TEAL, 3)
add_text(
    slide,
    "Landry Bougang Fotso  |  Founder & CEO  |  Kigali, Rwanda",
    0.76,
    5.67,
    8.8,
    0.35,
    size=11,
)
add_text(slide, "salanor.com", 0.76, 6.13, 4, 0.3, size=10, color=TEAL)
rect(slide, 9.1, 1.55, 3.35, 4.52, fill=SURFACE, line=LINE)
add_text(slide, "THINK", 9.48, 2.05, 2.6, 0.36, size=11, color=TEAL, bold=True)
add_text(slide, "Stoa", 9.48, 2.46, 2.6, 0.48, size=24, bold=True, font=FONT_DISPLAY)
add_text(slide, "Governed AI workspace\nand gateway", 9.48, 3.03, 2.55, 0.65, size=12, color=MUTED)
line(slide, 9.48, 3.85, 12.05, 3.88, LINE, 1)
add_text(slide, "ACT", 9.48, 4.18, 2.6, 0.36, size=11, color=TEAL, bold=True)
add_text(slide, "Aegis", 9.48, 4.59, 2.6, 0.48, size=24, bold=True, font=FONT_DISPLAY)
add_text(slide, "Policy, approval\nand signed evidence", 9.48, 5.16, 2.55, 0.65, size=12, color=MUTED)


# 2. Problem
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "The problem",
    "AI enters the enterprise through two doors",
    "Both create risk. Most control stacks cover only part of the journey.",
    2,
)
card(
    slide,
    0.65,
    2.45,
    5.85,
    3.65,
    "People use AI",
    "Staff paste contracts, claims, customer data, and internal documents into personal AI accounts.\n\nLeadership gets productivity, but no reliable control over data exposure, model choice, or retention.",
    accent=AMBER,
    number="01",
    body_size=14,
)
card(
    slide,
    6.83,
    2.45,
    5.85,
    3.65,
    "AI systems act",
    "Agents call APIs, update records, move money, and trigger workflows.\n\nWhen something goes wrong, logs rarely show who authorized the action, which policy applied, or whether a human approved it.",
    accent=RED,
    number="02",
    body_size=14,
)
add_text(
    slide,
    "Blocking AI does not solve adoption. Uncontrolled AI does not survive audit.",
    1.65,
    6.43,
    10.1,
    0.35,
    size=15,
    bold=True,
    align=PP_ALIGN.CENTER,
)


# 3. Why now
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "Why now",
    "Capability moved faster than enterprise control",
    "The gap is no longer theoretical. It appears in daily work, real workflows, and regulatory review.",
    3,
)
card(
    slide,
    0.65,
    2.48,
    3.85,
    3.4,
    "Model choice is fragmenting",
    "Employees and developers can reach many frontier models. One vendor policy cannot govern the whole estate.",
    number="01",
)
card(
    slide,
    4.74,
    2.48,
    3.85,
    3.4,
    "Actions are becoming autonomous",
    "AI is moving from drafting answers to operating tools. Control must run before the action, not after an incident.",
    number="02",
)
card(
    slide,
    8.83,
    2.48,
    3.85,
    3.4,
    "Evidence is becoming mandatory",
    "Regulated teams need logs, accountability, retention, and intervention. Screenshots and prompt history are not a control system.",
    number="03",
)
add_text(
    slide,
    "Design implication: governance must sit between identity, data, models, and tools.",
    0.68,
    6.28,
    11.9,
    0.4,
    size=15,
    color=TEAL,
    bold=True,
)
add_text(
    slide,
    "Sources: EU AI Act logging obligations; Anthropic Threat Intelligence Report, September 2026.",
    0.68,
    6.72,
    11.9,
    0.25,
    size=8,
    color=DIM,
)


# 4. Solution
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "The solution",
    "One control plane for how people think and systems act",
    "A familiar AI experience on top. Identity, policy, redaction, approval, and evidence underneath.",
    4,
)
rect(slide, 0.65, 2.45, 5.72, 3.65, fill=SURFACE, line=TEAL)
pill(slide, "Product 1 | Adopt now", 0.95, 2.77, 1.85)
add_text(slide, "Stoa", 0.95, 3.28, 2.7, 0.5, size=25, bold=True, font=FONT_DISPLAY)
add_text(
    slide,
    "A governed AI workspace and gateway for staff, apps, and approved models.",
    0.95,
    3.9,
    4.9,
    0.7,
    size=14,
    color=MUTED,
)
add_rich_text(
    slide,
    ["DLP before the provider", "Work identity and policy", "Audit and retention"],
    0.95,
    4.78,
    4.6,
    0.98,
    size=11,
    color=WHITE,
    bullet=True,
    gap=4,
)
rect(slide, 6.96, 2.45, 5.72, 3.65, fill=SURFACE, line=TEAL)
pill(slide, "Product 2 | Expand", 7.26, 2.77, 1.65)
add_text(slide, "Aegis", 7.26, 3.28, 2.7, 0.5, size=25, bold=True, font=FONT_DISPLAY)
add_text(
    slide,
    "Runtime governance for AI tools, APIs, payments, and operational workflows.",
    7.26,
    3.9,
    4.9,
    0.7,
    size=14,
    color=MUTED,
)
add_rich_text(
    slide,
    ["Allow, deny, or require approval", "Signed causal trace", "Exportable compliance evidence"],
    7.26,
    4.78,
    4.8,
    0.98,
    size=11,
    color=WHITE,
    bullet=True,
    gap=4,
)
arrow(slide, 6.42, 4.12, 0.48)
add_text(
    slide,
    "Same identity | same policy | same audit story",
    3.05,
    6.4,
    7.25,
    0.35,
    size=14,
    color=TEAL,
    bold=True,
    align=PP_ALIGN.CENTER,
)


# 5. Stoa
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "Product | Stoa",
    "The company-approved place to use AI",
    "ChatGPT-class usability with controls that enterprise buyers can explain.",
    5,
)
steps = [
    ("1", "Work identity", "Users sign in through the organisation, not personal AI accounts."),
    ("2", "Policy and DLP", "Prompts and files are scanned. Sensitive values are masked or blocked before the provider."),
    ("3", "Approved models", "One workspace routes to the models and contracts the organisation permits."),
    ("4", "Visible proof", "Users can see what the model received. Admins get metadata, audit, retention, and exports."),
]
for i, (num, title, body) in enumerate(steps):
    y = 2.35 + i * 1.06
    pill(slide, num, 0.75, y + 0.02, 0.42)
    add_text(slide, title, 1.42, y, 2.2, 0.32, size=14, bold=True)
    add_text(slide, body, 3.55, y, 8.8, 0.55, size=12, color=MUTED)
line(slide, 0.78, 6.63, 12.45, 6.66, LINE, 1)
add_text(
    slide,
    "Deployment: multi-tenant SaaS | dedicated tenant | private customer cloud",
    0.78,
    6.82,
    11.5,
    0.28,
    size=10,
    color=TEAL,
    bold=True,
)


# 6. Aegis
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "Product | Aegis",
    "Govern the action before it happens",
    "Policy enforcement and human approval in the path of sensitive operations.",
    6,
)
flow = [
    ("Agent", "Attempts a refund,\npayout, export, or\nadmin change"),
    ("Policy", "Evaluates identity,\ncontext, amount,\nand rule version"),
    ("Decision", "Allow | deny |\nrequire named\napproval"),
    ("Evidence", "Signed trace,\nverification, and\naudit export"),
]
for i, (title, body) in enumerate(flow):
    x = 0.72 + i * 3.16
    rect(slide, x, 2.62, 2.55, 2.62, fill=SURFACE)
    add_text(slide, f"0{i+1}", x + 0.24, 2.86, 0.42, 0.3, size=9, color=TEAL, bold=True)
    add_text(slide, title, x + 0.24, 3.32, 2.0, 0.38, size=16, bold=True)
    add_text(slide, body, x + 0.24, 3.93, 2.05, 0.9, size=12, color=MUTED)
    if i < len(flow) - 1:
        arrow(slide, x + 2.63, 3.72, 0.42)
add_text(
    slide,
    "Built for the moment AI stops being a chat and starts changing the world around it.",
    1.16,
    5.72,
    11.0,
    0.48,
    size=17,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "APS-1 signed events | policy p50 below 5ms target | Merkle witness | compliance bundles",
    1.2,
    6.34,
    10.9,
    0.3,
    size=10,
    color=TEAL,
    align=PP_ALIGN.CENTER,
)


# 7. Architecture
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "Platform",
    "Control stays close to the customer",
    "One trust zone for chat, gateway, policy, and audit. Customer-approved providers remain replaceable.",
    7,
)
layers = [
    (0.7, 2.42, 2.2, 2.65, "Identity", ["SSO / OIDC", "Users and roles", "Service accounts"]),
    (3.18, 2.42, 2.2, 2.65, "Control", ["Policy engine", "DLP / redaction", "Approval routing"]),
    (5.66, 2.42, 2.2, 2.65, "Experience", ["Stoa web", "AI gateway", "SDK / n8n"]),
    (8.14, 2.42, 2.2, 2.65, "Evidence", ["Signed events", "Trace and verify", "Retention / export"]),
    (10.62, 2.42, 2.0, 2.65, "Providers", ["OpenAI", "Anthropic", "Google / Azure"]),
]
for x, y, w, h, title, bullets in layers:
    rect(slide, x, y, w, h, fill=SURFACE)
    add_text(slide, title, x + 0.24, y + 0.3, w - 0.48, 0.35, size=14, bold=True)
    add_rich_text(
        slide,
        bullets,
        x + 0.24,
        y + 0.96,
        w - 0.48,
        1.35,
        size=10,
        color=MUTED,
        bullet=True,
        gap=7,
    )
add_text(
    slide,
    "SaaS",
    1.12,
    5.6,
    1.2,
    0.3,
    size=11,
    color=TEAL,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "Dedicated",
    4.03,
    5.6,
    1.4,
    0.3,
    size=11,
    color=TEAL,
    bold=True,
    align=PP_ALIGN.CENTER,
)
add_text(
    slide,
    "Private cloud",
    7.16,
    5.6,
    1.65,
    0.3,
    size=11,
    color=TEAL,
    bold=True,
    align=PP_ALIGN.CENTER,
)
line(slide, 1.0, 6.13, 11.9, 6.16, LINE, 1)
add_text(
    slide,
    "Customer BYOK | metadata-first logging | configurable retention | no source-code transfer required",
    1.1,
    6.34,
    11.2,
    0.3,
    size=11,
    color=MUTED,
    align=PP_ALIGN.CENTER,
)


# 8. Concrete use case
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "Example",
    "One customer workflow, governed end to end",
    "Support work starts in Stoa. A refund action is governed by Aegis.",
    8,
)
stages = [
    ("01", "Upload", "Invoice or complaint enters Stoa."),
    ("02", "Protect", "Names, IDs, email, phone, and account data are redacted."),
    ("03", "Reason", "Approved model receives only the protected payload."),
    ("04", "Act", "Agent proposes a USD 2,500 refund."),
    ("05", "Control", "Policy pauses the call above USD 1,000."),
    ("06", "Prove", "Named approval and signed trace are exported."),
]
for i, (num, title, body) in enumerate(stages):
    row = 0 if i < 3 else 1
    col = i if i < 3 else i - 3
    x = 0.72 + col * 4.18
    y = 2.33 + row * 2.05
    rect(slide, x, y, 3.72, 1.6, fill=SURFACE)
    add_text(slide, num, x + 0.22, y + 0.22, 0.4, 0.25, size=9, color=TEAL, bold=True)
    add_text(slide, title, x + 0.72, y + 0.18, 2.55, 0.35, size=14, bold=True)
    add_text(slide, body, x + 0.22, y + 0.76, 3.28, 0.58, size=11, color=MUTED)
add_text(
    slide,
    "Result: familiar AI productivity without losing control of secrets, money, or accountability.",
    0.9,
    6.48,
    11.55,
    0.36,
    size=15,
    color=TEAL,
    bold=True,
    align=PP_ALIGN.CENTER,
)


# 9. ICP / GTM
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "Go to market",
    "Start where AI risk is already budgeted",
    "Sell one measurable control problem, then expand across the organisation.",
    9,
)
card(
    slide,
    0.65,
    2.42,
    3.75,
    3.78,
    "Beachhead",
    "Regulated teams with active AI use:\n\n• Fintech and payments\n• Insurance operations\n• Law, audit, and consulting\n• Bank innovation teams",
    number="01",
    body_size=12,
)
card(
    slide,
    4.79,
    2.42,
    3.75,
    3.78,
    "Land",
    "A paid 4 to 8 week pilot on one workflow:\n\n• Sensitive AI workspace\n• Refund or payout gate\n• KYC escalation\n• Customer-data export",
    number="02",
    body_size=12,
)
card(
    slide,
    8.93,
    2.42,
    3.75,
    3.78,
    "Expand",
    "From one team to the control plane:\n\n• More users and models\n• Gateway for internal apps\n• Aegis on sensitive actions\n• Dedicated or private deploy",
    number="03",
    body_size=12,
)
add_text(
    slide,
    "Channel leverage: implementation partners, legal/compliance firms, and regional cloud partners.",
    0.68,
    6.46,
    11.9,
    0.32,
    size=12,
    color=TEAL,
    bold=True,
)


# 10. Business model
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "Business model",
    "Recurring software revenue with paid implementation",
    "Customers pay Salanor for control and evidence. Enterprise model usage stays BYOK.",
    10,
)
card(
    slide,
    0.65,
    2.42,
    3.7,
    3.62,
    "Pilot",
    "USD 2k to 8k implementation\n\nOne workflow or one team\n4 to 8 weeks\nDefined success criteria",
    number="01",
    body_size=13,
)
card(
    slide,
    4.82,
    2.42,
    3.7,
    3.62,
    "Platform",
    "Aegis Team from USD 299/month\n\nStoa seats or user bands\nAnnual enterprise contracts\nSupport and retention tiers",
    number="02",
    body_size=13,
)
card(
    slide,
    8.99,
    2.42,
    3.7,
    3.62,
    "Enterprise",
    "Setup + annual platform + committed user band\n\nDedicated or private deployment\nSLA, SSO, SIEM, policy packs\nCustomer-owned model contracts",
    number="03",
    body_size=12,
)
add_text(
    slide,
    "Expansion engine: workspace users lead to gateway volume and governed actions.",
    0.68,
    6.4,
    11.8,
    0.32,
    size=14,
    color=TEAL,
    bold=True,
)


# 11. Positioning
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "Positioning",
    "The missing layer between AI access and AI action",
    "Salanor does not replace models, identity, SIEM, or workflow engines. It connects their controls.",
    11,
)
rows = [
    ("Consumer AI", "Excellent user experience", "No organisation-specific control or proof"),
    ("Enterprise copilots", "Suite-native governance", "Bound to one ecosystem"),
    ("AI gateways / DLP", "Protect model requests", "Limited action-level governance"),
    ("Observability", "Helps engineers debug", "Logs are not signed policy evidence"),
    ("Salanor", "Workspace + gateway + action control", "One policy and evidence plane"),
]
ys = [2.38, 3.2, 4.02, 4.84, 5.66]
for idx, ((name, strength, gap), y) in enumerate(zip(rows, ys)):
    fill = TEAL_DARK if idx == len(rows) - 1 else SURFACE
    border = TEAL if idx == len(rows) - 1 else LINE
    rect(slide, 0.68, y, 11.95, 0.64, fill=fill, line=border)
    add_text(slide, name, 0.93, y + 0.15, 2.25, 0.28, size=11, bold=True)
    add_text(slide, strength, 3.35, y + 0.15, 3.55, 0.28, size=11, color=MUTED)
    add_text(slide, gap, 7.25, y + 0.15, 4.95, 0.28, size=11, color=WHITE if idx == len(rows) - 1 else MUTED)
add_text(slide, "Category", 0.93, 2.06, 2.0, 0.22, size=8, color=DIM, bold=True)
add_text(slide, "What it does well", 3.35, 2.06, 3.0, 0.22, size=8, color=DIM, bold=True)
add_text(slide, "Control gap", 7.25, 2.06, 3.0, 0.22, size=8, color=DIM, bold=True)


# 12. Progress
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "Progress",
    "The platform exists. The next proof is commercial.",
    "Early access product, working end-to-end demos, and a clear design-partner motion.",
    12,
)
add_text(slide, "BUILT", 0.72, 2.34, 2.1, 0.32, size=10, color=TEAL, bold=True)
add_rich_text(
    slide,
    [
        "Stoa workspace: multi-model chat, file protection, DLP, admin, audit",
        "Aegis: policy, approvals, signed traces, verification, compliance export",
        "Shared identity, organisation isolation, provider configuration, retention",
        "SDK and reference workflow for a governed support-to-refund scenario",
    ],
    0.72,
    2.83,
    5.75,
    2.8,
    size=12,
    color=WHITE,
    bullet=True,
    gap=13,
)
line(slide, 6.65, 2.32, 6.68, 6.25, LINE, 1)
add_text(slide, "NEXT VALIDATION", 7.05, 2.34, 3.4, 0.32, size=10, color=AMBER, bold=True)
add_rich_text(
    slide,
    [
        "First paid design partner on one high-risk workflow",
        "Security and procurement pack tested with a regulated buyer",
        "Reference case with measurable control and audit outcomes",
        "Repeatable dedicated and private deployment package",
    ],
    7.05,
    2.83,
    5.2,
    2.8,
    size=12,
    color=WHITE,
    bullet=True,
    gap=13,
)
add_text(
    slide,
    "Stage: early access | No customer or revenue numbers are claimed in this deck.",
    0.75,
    6.5,
    11.7,
    0.28,
    size=10,
    color=MUTED,
)


# 13. Roadmap
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "Execution",
    "A focused 12-month path",
    "Commercial validation first. Enterprise breadth follows customer pull.",
    13,
)
milestones = [
    ("0 to 3 months", "Validate", ["2 paid pilots", "One reference workflow", "Buyer-approved success criteria"]),
    ("3 to 6 months", "Package", ["SSO hardening", "Dedicated deployment", "Security and DPA pack"]),
    ("6 to 12 months", "Expand", ["Gateway adoption", "Aegis action coverage", "Channel-led pipeline"]),
]
for i, (period, title, bullets) in enumerate(milestones):
    x = 0.72 + i * 4.18
    rect(slide, x, 2.48, 3.72, 3.52, fill=SURFACE)
    pill(slide, period, x + 0.25, 2.76, 1.35)
    add_text(slide, title, x + 0.25, 3.35, 2.8, 0.4, size=20, bold=True, font=FONT_DISPLAY)
    add_rich_text(
        slide,
        bullets,
        x + 0.25,
        4.13,
        3.12,
        1.35,
        size=11,
        color=MUTED,
        bullet=True,
        gap=10,
    )
add_text(
    slide,
    "Operating rule: do not build enterprise complexity before a buyer requires it in writing.",
    0.75,
    6.4,
    11.7,
    0.35,
    size=14,
    color=TEAL,
    bold=True,
)


# 14. Team and ask
slide = prs.slides.add_slide(blank)
add_bg(slide)
slide_title(
    slide,
    "Team and ask",
    "Building the control layer before autonomy becomes normal",
    None,
    14,
)
rect(slide, 0.72, 2.15, 5.35, 3.92, fill=SURFACE)
pill(slide, "Founder", 1.02, 2.48, 0.9)
add_text(
    slide,
    "Landry Bougang Fotso",
    1.02,
    3.04,
    4.4,
    0.45,
    size=22,
    bold=True,
    font=FONT_DISPLAY,
)
add_text(slide, "Founder & CEO | Kigali, Rwanda", 1.02, 3.57, 4.4, 0.34, size=12, color=TEAL)
add_text(
    slide,
    "Product and engineering founder building infrastructure for high-stakes AI decisions. Previously built Aether through Rwanda's Plug-In Ventures before focusing Salanor on AI governance.",
    1.02,
    4.18,
    4.45,
    1.2,
    size=12,
    color=MUTED,
)
add_text(slide, "WHAT WE ARE LOOKING FOR", 6.75, 2.28, 4.6, 0.3, size=10, color=TEAL, bold=True)
add_rich_text(
    slide,
    [
        "Two regulated design partners with a real AI workflow",
        "Introductions to CTO, CISO, risk, and compliance leaders",
        "Pre-seed conversations with investors who understand enterprise infrastructure",
        "Partners for SSO, cloud deployment, and legal/compliance delivery",
    ],
    6.75,
    2.9,
    5.45,
    2.45,
    size=13,
    color=WHITE,
    bullet=True,
    gap=15,
)
line(slide, 6.75, 5.72, 12.1, 5.75, LINE, 1)
add_text(slide, "landry@salanor.com", 6.75, 6.0, 3.8, 0.3, size=12, color=TEAL, bold=True)
add_text(slide, "salanor.com", 10.2, 6.0, 1.9, 0.3, size=12, color=WHITE, align=PP_ALIGN.RIGHT)
add_text(
    slide,
    "Financing terms intentionally omitted until round size and milestones are fixed.",
    6.75,
    6.48,
    5.45,
    0.25,
    size=8,
    color=DIM,
)


# Core properties
prs.core_properties.title = "Salanor Startup Pitch Deck"
prs.core_properties.subject = "Control and proof for enterprise AI"
prs.core_properties.author = "Salanor Ltd"
prs.core_properties.company = "Salanor Ltd"
prs.core_properties.comments = (
    "Generated from the Salanor product and commercial documentation. "
    "Claims are intentionally limited to capabilities and current early-access status."
)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"Wrote {OUT} ({len(prs.slides)} slides)")
